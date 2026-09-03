"""Unit test for Case 2 (SecureBorder Solutions B.V.) deterministic pipeline."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import yaml

from aegis_phase1.v2.loader.case_profile import CaseProfileLoader
from aegis_phase1.v2.loader.preproc_catalog import PreprocCatalogLoader
from aegis_phase1.v2.orchestrator import Phase1Orchestrator


def test_case2_deterministic_pipeline(tmp_path: Path) -> None:
    case_dir = Path("cases/case2-secureborder")
    assert case_dir.exists(), "Case 2 directory missing"

    # 1. Test profile loader
    profile = CaseProfileLoader(case_dir).load()
    assert profile.company.name == "SecureBorder Solutions B.V."
    assert profile.company.scale == "LARGE"
    assert set(profile.applicable_regs) == {"AI_Act", "CRA", "GDPR", "NIS2"}
    assert len(profile.declaration_gaps) == 0
    assert len(profile.architecture.systems) >= 12
    assert len(profile.architecture.data_stores) >= 7

    # 2. Test orchestrator deterministic run
    orch = Phase1Orchestrator(
        work_dir=str(tmp_path / "work"),
        preproc_catalog=PreprocCatalogLoader(preproc_root="preproc_out"),
        case_profile_loader=CaseProfileLoader(case_dir),
    )
    orch.load(str(case_dir), "Methodology-main/00_METHODOLOGY/PREPROCESSING")
    out_dir = tmp_path / "out"
    orch.generate_deterministic_docs(str(out_dir))
    written = orch.state.get("output_paths", {})

    expected_docs = [
        "AEGIS-P1-04",
        "AEGIS-P1-05",
        "AEGIS-P1-06",
        "AEGIS-P1-07",
        "AEGIS-P1-07b",
        "AEGIS-P1-XLSX",
    ]
    for doc in expected_docs:
        assert doc in written, f"Expected {doc} in output_paths"
        path = Path(written[doc])
        assert path.exists() and path.stat().st_size > 500, f"Artifact {doc} missing or empty"

    # 3. Verify Doc 04
    doc04_text = Path(written["AEGIS-P1-04"]).read_text(encoding="utf-8")
    assert "SecureBorder Solutions B.V." in doc04_text
    assert "LARGE" in doc04_text
    assert "€120M" in doc04_text

    # 4. Verify Doc 07 frontmatter and matrix
    doc07_text = Path(written["AEGIS-P1-07"]).read_text(encoding="utf-8")
    parts = doc07_text.split("---", 2)
    fm = yaml.safe_load(parts[1])
    assert fm["case_study"] == "SecureBorder Solutions B.V."
    assert fm["applicable_regs"] == ["AI_Act", "CRA", "GDPR", "NIS2"]
    assert "✅ AI_Act" in doc07_text
    assert "✅ CRA" in doc07_text
    assert "✅ GDPR" in doc07_text
    assert "✅ NIS2" in doc07_text

    # 5. Verify Excel workbook
    xlsx_path = Path(written["AEGIS-P1-XLSX"])
    assert xlsx_path.name == "Case_02_Phase1.xlsx"
    wb = openpyxl.load_workbook(str(xlsx_path))
    assert "COVER" in wb.sheetnames
    assert "COMPANY" in wb.sheetnames
    assert "REGULATIONS" in wb.sheetnames
    assert "COVERAGE" in wb.sheetnames
